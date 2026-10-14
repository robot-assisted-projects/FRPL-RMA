"""Exact original 2024–25 San Diego Unified source/cohort audit; no fits or imports."""
import argparse
import ast
from collections import Counter, defaultdict
from datetime import date, datetime
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import re
import statistics
import subprocess
import types
import zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'data/source/san-diego-district-audit.json'
LEA, NATIVE_LEA, CDS_DISTRICT = '0634320', 'CA-3768338', '3768338'
YEAR, TEST_YEAR = '2024-2025', 2025
OPERATIONAL = {'1', '3', '4', '5', '8'}
GRADES = ['PK', 'KG'] + [str(g) for g in range(1, 14)] + ['UG', 'AE']
SCHOOL_TYPES = {'7', '9', '10'}
SUBJECTS = {'math': '2', 'reading': '1'}
MIN_VALID = 11
SOURCES = {'directory': {'bytes': 13352819, 'path': 'data/raw/ccd_sch_029_2425_w_1a_073025.zip', 'sha256': '39326da788aa322353d20ceaf8ad4baed26272502cd05b066cf6c594988b21ab', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_sch_029_2425_w_1a_073025.zip'}, 'membership': {'bytes': 212696691, 'path': 'data/raw/ccd_sch_052_2425_l_1a_073025.zip', 'sha256': '4a7f660c5fc5eaae488dd02fd43498f349fc828b227edd0970d5b6995ead4d4d', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_sch_052_2425_l_1a_073025.zip'}, 'lea_directory': {'bytes': 2830266, 'path': 'data/raw/ccd_lea_029_2425_w_1a_073025.zip', 'sha256': '2745169e4bc7cd53adff179830c3d4b9bb035ac79580ecea01ef389a161936a9', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_lea_029_2425_w_1a_073025.zip'}, 'lea_membership': {'bytes': 68497561, 'path': 'data/raw/ccd_lea_052_2425_l_1a_073025.zip', 'sha256': '501d72720a01c26e0e041cd3b1aa6653d0a94725ba0a1e85c42c4f183bc627ba', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_lea_052_2425_l_1a_073025.zip'}, 'ccd_state_notes': {'bytes': 210984, 'path': 'data/raw/SY_2024-25_CCD_Final_2a_Data_Notes.xlsx', 'sha256': '6476b6b28b700389e261b87377d9cd249935503d53b5919fe1b7ef8b1052d1a7', 'url': 'https://nces.ed.gov/ccd/xls/SY_2024-25_CCD_Final_2a_Data_Notes.xlsx'}, 'assessment': {'path': 'data/raw/california/sb_ca2025_1_csv_v1.zip', 'url': 'https://caaspp-elpac.ets.org/caaspp/researchfiles/sb_ca2025_1_csv_v1.zip', 'bytes': 7789580, 'sha256': 'b1fea7cc714141ffe9c67ac91bad3eaa24282c49d37a18bd0f9ad6e2dc3b6e6b'}, 'frpm': {'path': 'data/raw/california/frpm2425.xlsx', 'url': 'https://www.cde.ca.gov/Ds/ad/documents/frpm2425.xlsx', 'bytes': 2185842, 'sha256': '86fb8bc64ce0e433468f4a4b9038115a82031fee56125f76bacd56037a654a40'}, 'research_layout': {'path': 'data/raw/california/2025_SBAC_Research_File_Layout.xlsx', 'url': 'https://caaspp-elpac.ets.org/caaspp/docs/2025_SBAC_Research%20File%20Layout.xlsx', 'bytes': 34870, 'sha256': 'a1c6bfb0c88333413ad612b0aa47a7ad08a85c6d7cb1c27f7ff9facecbf6ab2e'}, 'research_page': {'path': 'data/raw/california/research-page.html', 'url': 'https://caaspp-elpac.ets.org/caaspp/ResearchFileListSB?lstCounty=00&lstDistrict=00000&lstTestType=B&lstTestYear=2025&ps=true', 'bytes': 62338, 'sha256': 'c2d784a0e685716b1d7377f1333d9e9da2457de76ef024ebf1a4faa4ac52ffc3'}}
POLICY = {
    'identity': 'Exact original CCD LEA 0634320 / CA-3768338 with full CA-3768338-sssssss school IDs mapped to full 3768338sssssss CDS. Names and suffix-only matches never join.',
    'year': 'Only 2024–25 CCD/Census Day eligibility and 2025 Smarter Balanced outcomes.',
    'income': 'Own same-year CALPADS Census Day K–12 FRPM eligible count divided by own K–12 enrollment, including applications, alternate income forms, direct and categorical certification. Eligibility differs from universally available meals, age5–17 measures and LCFF unduplicated pupil count.',
    'assessment': 'Smarter Balanced B, All Students group1, separate Math2/ELA1. Complete native expected grades3–8 use verified Overall Total valid-performance-score counts as aggregation weights. Grade13 all grades3–8 and11 is retained and reconciled; it never recovers suppressed/missing component grades. CAA and other assessments never pool.',
    'counts': 'Overall Total equals the sum of four performance levels and Total Students Tested with Scores; met-plus-exceeded counts reconcile to published two-decimal rates. At least11 valid scores per expected grade. Enrollment and tested participation never replace valid scores.',
    'scope': 'Prospective lower profiles require agreeing CCD/native lower configuration and applicable grades3–8; substantive lower/high conflict is separate. P versus KG low-bound differences are retained with California preschool-reporting notes, without changing tested grades. CCD missing/suppressed membership never means zero or replaces the native income denominator.',
    'charters': 'Retain exact roster charter flags. Reporting-prefix charters outside exact CCD LEA membership remain outside; no name, prefix or funding label grants district membership.',
    'programs': 'Only original administrative school/educational-option/charter labels are verified; magnet/selective/virtual-provider/admissions attribution remains unknown. A name is not program evidence.',
    'combined': 'Equal mean of usable Math and ELA proficiency; retain separate subject valid-score denominators, never a synthetic paired tested count.',
    'approval': 'Source/cohort audit only; historical source/model approvals remain false and numerical/normalized review is separate.',
}
RAW_INPUT_SHA = '5db2e5058ea6243eb5843cf7f5e8d55b583551576938b7a95b3ea77331f9ad3e'
ORIGINAL_SCHEMA_SHA = '0720d1f22dde683426ffe2b766b3bf48400c816c558a6cd75c5d9fadc6070619'
DEFINITION_INPUT_SHA = '8617ecb470fb3e7e5b288a1a504e1d86af060333d02ef373ff37deb48fcb9696'
ORIGINAL_IDENTITY_SHA = '488a7bc2a9ec19670e87bc6e1071813471b22cd8146e41786ab20024d3f96bd2'
POLICY_SHA = 'ca6b2681810ce6f80cf79092b61c325a8abb0209f33a6ce5b9b42440c8ccfb53'
RUNTIME_CONTRACT_SHA = '2b725ef1988e3944d25a5c905be3d8808cf8f6ca494e3de5972c60ffec8807c6'
ORIGINAL_INVENTORIES = {'assessment': 1816, 'directory': 175, 'entities': 216, 'frpm': 217, 'lea_directory': 1, 'lea_membership': 16, 'membership': 1348}
ORIGINAL_METADATA = {'assessment': 'b7632e1f078a6eb54e7c576c916cef715a0b3b5fa3719349e35578a131870799', 'directory': '1d65a2bb1bed90acbf1f282b4825bcfe9f72aa31eb29a84e2fb1996f9ace0983', 'entities': '632d719cadcce9d050686b0b9db5fb13c13ac6023b064273029dd803fa779510', 'frpm': '37cad732117d753cf51ec141c78aa7d6e171c2086aca3119ea07c640d807b583', 'lea_directory': '9a26c257086ce04ee9aa5d185df5c5318a51a581a5015c2e8ed071fec971d7de', 'lea_membership': '68f5560e90ae5abd35b04750c4b7ebd95d37217a41d7f79f928ba4dc3e7764d2', 'membership': '230663b60efff9cafa9eb3b26bcebf2241b18781ee64b04174ec26a0ceb368d8'}

def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(',', ':'), ensure_ascii=False, allow_nan=False).encode()).hexdigest()

def native_cell(value):
    return dict(cell_type=type(value).__name__, iso8601=value.isoformat()) if isinstance(value, (date, datetime)) else value

def code_contract(code):
    def constant(value):
        if isinstance(value, types.CodeType): return code_contract(value)
        if isinstance(value, tuple): return ['tuple', [constant(v) for v in value]]
        if isinstance(value, frozenset): return ['frozenset', sorted((constant(v) for v in value), key=repr)]
        if isinstance(value, slice): return ['slice', constant(value.start), constant(value.stop), constant(value.step)]
        if isinstance(value, bytes): return ['bytes', value.hex()]
        if value is Ellipsis: return ['ellipsis']
        return [type(value).__name__, value]
    return dict(argcount=code.co_argcount, posonly=code.co_posonlyargcount, kwonly=code.co_kwonlyargcount,
        names=list(code.co_names), varnames=list(code.co_varnames), freevars=list(code.co_freevars),
        cellvars=list(code.co_cellvars), flags=code.co_flags, bytecode=code.co_code.hex(), constants=[constant(v) for v in code.co_consts])

def semantic_sources():
    names = ['cds','ccd_identity','count','ccd_count','grade','offered_evidence',
             'enrollment_evidence','income_evidence','valid_grade','subject_evidence',
             'summarize','validate_definition_references','original_identity_fingerprint']
    text=Path(__file__).read_text()
    tree=ast.parse(text)
    nodes={node.name:node for node in tree.body if isinstance(node,ast.FunctionDef)and node.name in names}
    if set(nodes)!=set(names):raise ValueError('Complete original semantic function declarations required')
    return {name:ast.get_source_segment(text,nodes[name])for name in names}

def require_live_semantics():
    # Compile without executing the current script, after its original semantic
    # declarations are literal-verified. Whole-module compilation preserves the
    # interpreter's import-aware optimizations and avoids a Python-version pin.
    compiled=compile(Path(__file__).read_text(),str(__file__),'exec',dont_inherit=True)
    reference={c.co_name:c for c in compiled.co_consts if isinstance(c,types.CodeType)}
    for name in semantic_sources():
        actual=globals()[name]
        if not isinstance(actual,types.FunctionType)or code_contract(actual.__code__)!=code_contract(reference[name]):
            raise ValueError('Current original semantic function changed: '+name)

def runtime_contract():
    return dict(lea=LEA, native_lea=NATIVE_LEA, cds_district=CDS_DISTRICT, year=YEAR, test_year=TEST_YEAR,
        operational=sorted(OPERATIONAL), grades=GRADES, school_types=sorted(SCHOOL_TYPES), subjects=SUBJECTS,
        min_valid=MIN_VALID, sources=SOURCES, policy=POLICY, raw=RAW_INPUT_SHA, schema=ORIGINAL_SCHEMA_SHA,
        definitions=DEFINITION_INPUT_SHA, identity=ORIGINAL_IDENTITY_SHA,
        inventories=ORIGINAL_INVENTORIES, metadata=ORIGINAL_METADATA,
        semantic_sources=semantic_sources())

def require_runtime_contract():
    if (POLICY_SHA != 'ca6b2681810ce6f80cf79092b61c325a8abb0209f33a6ce5b9b42440c8ccfb53' or RUNTIME_CONTRACT_SHA != '2b725ef1988e3944d25a5c905be3d8808cf8f6ca494e3de5972c60ffec8807c6'
        or fingerprint(POLICY) != 'ca6b2681810ce6f80cf79092b61c325a8abb0209f33a6ce5b9b42440c8ccfb53' or fingerprint(runtime_contract()) != '2b725ef1988e3944d25a5c905be3d8808cf8f6ca494e3de5972c60ffec8807c6'):
        raise ValueError('San Diego current immutable source policy/identity/schema/semantic runtime changed')
    require_live_semantics()

def require_current_originals():
    require_runtime_contract()
    for key, pin in SOURCES.items():
        path = ROOT / pin['path']
        if path.stat().st_size != pin['bytes']: raise ValueError('San Diego whole original size changed: '+key)
        h = hashlib.sha256()
        with path.open('rb') as f:
            for b in iter(lambda:f.read(1048576),b''): h.update(b)
        if h.hexdigest() != pin['sha256']: raise ValueError('San Diego whole original fingerprint changed: '+key)

def cds(raw):
    parts = [raw.get(k) for k in ['County Code', 'District Code', 'School Code']]
    if any(type(v) is not str or not re.fullmatch(r'\d{'+str(size)+'}', v) for v,size in zip(parts,[2,5,7])):
        raise ValueError('Full original county/district/school CDS components required')
    if ''.join(parts[:2]) != CDS_DISTRICT: raise ValueError('Original native CDS district changed')
    return ''.join(parts)

def ccd_identity(raw, school=True):
    if (raw.get('SCHOOL_YEAR'),raw.get('ST'),raw.get('FIPST'),raw.get('STATENAME'),raw.get('LEAID'),raw.get('ST_LEAID')) != (YEAR,'CA','06','CALIFORNIA',LEA,NATIVE_LEA):
        raise ValueError('Same-year complete San Diego CCD agency identity required')
    if not school: return LEA
    key, native = raw.get('NCESSCH'),raw.get('ST_SCHID')
    if type(key) is not str or not re.fullmatch(LEA+r'\d{5}',key) or type(native)is not str or not re.fullmatch(NATIVE_LEA+r'-\d{7}',native):
        raise ValueError('Exact complete CCD school/native identity required')
    return key, native[3:].replace('-', '')

def count(raw):
    if raw is None or raw in ('','*','N/A'): return None
    if type(raw) is bool or type(raw) not in (str,int,float): raise ValueError('Native counts must be numeric, never booleans')
    value=float(raw)
    if not math.isfinite(value) or value<0 or int(value)!=value: raise ValueError('Invalid nonnegative integer native count')
    return int(value)

def ccd_count(raw):
    if not raw or raw.get('DMS_FLAG') not in {'Reported','Derived'} or not re.fullmatch(r'\d+',raw.get('STUDENT_COUNT','')): return None
    return int(raw['STUDENT_COUNT'])

def grade(raw):
    if raw in ('P','PK'): return -1
    if raw in ('K','KG'): return 0
    return int(raw) if type(raw)is str and raw.isdigit() and 0<=int(raw)<=12 else None

def offered_evidence(raw, profile):
    if raw['IGOFFERED']!='As reported' or raw['NOGRADES']!='No' or any(raw['G_'+g+'_OFFERED']not in {'Yes','No'}for g in GRADES): raise ValueError('Explicit original grades-offered flags required')
    ccd_lower=any(raw['G_'+g+'_OFFERED']=='Yes'for g in ['PK','KG']+[str(n)for n in range(1,9)])
    ccd_high=any(raw['G_'+str(n)+'_OFFERED']=='Yes'for n in range(9,13))
    ccd_scope='other'if any(raw['G_'+g+'_OFFERED']=='Yes'for g in ['13','UG','AE'])else'mixed'if ccd_lower and ccd_high else'ES'if ccd_lower else'HS'if ccd_high else'other'
    low,high=grade(profile['Low Grade']),grade(profile['High Grade'])
    span_valid=low is not None and high is not None and low<=high and profile['School Type']!='N/A'
    native_scope='unknown'if not span_valid else'ES'if high<9 else'HS'if low>=9 else'mixed'
    expected=list(range(max(low,3),min(high,8)+1))if span_valid else[]
    ccd_tested=[n for n in range(3,9)if raw['G_'+str(n)+'_OFFERED']=='Yes']
    agreement=ccd_scope==native_scope and (native_scope!='ES' or expected==ccd_tested)
    return dict(ccd_scope=ccd_scope,native_scope=native_scope,configuration_agreement=agreement,
        ccd_raw_low=raw['GSLO'],ccd_raw_high=raw['GSHI'],native_raw_low=profile['Low Grade'],native_raw_high=profile['High Grade'],
        literal_bounds_agree=(grade(raw['GSLO']),grade(raw['GSHI']))==(low,high),
        expected_tested_grades=expected,ccd_offered_tested_grades=ccd_tested,
        pure_lower_candidate=agreement and native_scope=='ES' and bool(expected),
        lower_bound_reporting_note='California does not report full preschool enrollment in CCD/CALPADS. P/KG raw differences remain explicit; tested grades are independently compared.',
        exclusion=None if agreement else'CCD/native grade configuration conflict; retain outside prospective pure scope')

def enrollment_evidence(total, grades, offered):
    values={g:dict(source_row=e['source_row'],raw_count=e['raw']['STUDENT_COUNT'],dms_flag=e['raw']['DMS_FLAG'],count=ccd_count(e['raw']))for g,e in grades.items()}
    known=[v['count']for g,v in values.items()if g!='Not Specified'and v['count']is not None]
    high=[values.get('Grade '+str(n))for n in range(9,13)]
    unknown_high=any(v is not None and v['count']is None for v in high)
    high_positive=any(v is not None and v['count']is not None and v['count']>0 for v in high)
    lower_positive=any(values.get('Grade '+str(n),{}).get('count',0)>0 for n in range(3,9)if values.get('Grade '+str(n),{}).get('count')is not None)
    return dict(total_source_row=total['source_row']if total else None,total_raw=total['raw']if total else None,
        total=ccd_count(total['raw'])if total else None,grade_records=values,known_grade_sum=sum(known),
        complete_published_grade_sum=all(v['count']is not None for v in values.values()),
        high_positive=high_positive,high_unknown=unknown_high,known_tested_grade_enrollment_positive=lower_positive,
        note='Absent CCD grade rows remain absent; masked counts remain unknown. These administrative counts never supply native FRPM or valid-score denominators.')

def income_evidence(entry):
    raw=entry['raw']
    if raw['Academic Year']!=YEAR: raise ValueError('Same-year FRPM required')
    enrollment,eligible=count(raw['Enrollment (K-12)']),count(raw['FRPM Count (K-12)'])
    fraction=raw['Percent (%) Eligible FRPM (K-12)']
    if fraction in (None,'','*'): fraction=None
    elif type(fraction) not in (int,float)or type(fraction)is bool or not math.isfinite(fraction)or not 0<=fraction<=1: raise ValueError('Original FRPM fraction must be finite numeric0–1')
    if enrollment is not None and eligible is not None and eligible>enrollment: raise ValueError('FRPM exceeds own enrollment')
    reason='Missing native K–12 eligibility/enrollment/fraction'if None in (enrollment,eligible,fraction)else'Zero native K–12 enrollment'if enrollment==0 else'Native Fall1 certification unavailable'if raw['California Longitudinal Pupil Achievement Data System (CALPADS) Fall 1 Certification Status']!='Y'else None
    if reason is None and abs(eligible/enrollment-fraction)>1e-8: raise ValueError('Native FRPM fraction/counts mismatch')
    return dict(source_id='frpm',source_row=entry['source_row'],academic_year=YEAR,snapshot_date='2024-10-02',
        population='K–12 primary or short-term enrollment on CALPADS Census Day',enrollment=enrollment,eligible=eligible,
        published_fraction=fraction,percentage=None if reason else 100*eligible/enrollment,exclusion=reason,
        age5_17_counts_retained_separately=True,definition=POLICY['income'])

def valid_grade(entry):
    raw=entry['raw']; n=count(raw['Overall Total']); scores=count(raw['Total Students Tested with Scores']);met=count(raw['Count Standard Met and Above']);tested=count(raw['Total Students Tested'])
    text=raw['Percentage Standard Met and Above']
    if None in (n,scores,met,tested)or text in ('','*'): return None,'Missing or suppressed original valid-score/count/rate cell'
    if n<MIN_VALID: return None,'Fewer than11 valid performance scores'
    levels=[count(raw[k])for k in ['Count Standard Exceeded','Count Standard Met','Count Standard Nearly Met','Count Standard Not Met']]
    if None in levels: return None,'Missing or suppressed original performance-level count'
    rate=float(text)
    if (not math.isfinite(rate)or not 0<=rate<=100 or n!=scores or tested<n or met>n
        or sum(levels)!=n or sum(levels[:2])!=met or abs(rate-100*met/n)>.0051): raise ValueError('Original valid-score/performance counts and published percentage fail reconciliation')
    return dict(valid_scores=n,proficient=met,published_percentage=rate,percentage=100*met/n),None

def subject_evidence(records,expected,income):
    by_grade={}
    for entry in records:
        g=int(entry['raw']['Grade'])
        if g in by_grade: raise ValueError('Duplicate original school/subject/grade')
        by_grade[g]=entry
    applicable=bool(expected);reason=None if applicable else'No applicable tested grade in this level'
    if applicable and not set(expected)<=set(by_grade): reason='Missing an expected original tested grade'
    values=[]
    if reason is None:
        for g in expected:
            value,issue=valid_grade(by_grade[g])
            if issue: reason=issue;break
            values.append(value)
    if reason is None and income['exclusion']:reason=income['exclusion']
    n=sum(v['valid_scores']for v in values)if reason is None else None
    met=sum(v['proficient']for v in values)if reason is None else None
    whole=by_grade.get(13);whole_value,whole_reason=valid_grade(whole)if whole else(None,'No original all-grades13 row')
    whole_reconciles=(whole_value['valid_scores']==n and whole_value['proficient']==met)if whole_value is not None and n is not None else None
    return dict(applicable=applicable,source_usable=reason is None,expected_grades=expected,
        proficiency=None if reason else 100*met/n,valid_scores=n,proficient_count=met,exclusion=reason,
        exact_original_expected_grade_refs=[dict(source_id='assessment',source_row=by_grade[g]['source_row'],grade=g,raw=by_grade[g]['raw'])for g in expected if g in by_grade],
        all_original_school_subject_source_rows=[e['source_row']for e in records],
        native_all_grades13=None if whole is None else dict(source_row=whole['source_row'],raw=whole['raw'],verified=whole_value,exclusion=whole_reason),
        all_grades13_counts_match_expected_complete_components=whole_reconciles,
        missing_or_suppressed_components_never_recovered_from13=True)

def summarize(links,predicate):
    rows=[r for r in links if predicate(r)];subjects={}
    for subject in ['math','reading','combined']:
        eligible=[r for r in rows if r['subjects']['ES'][subject]['source_usable']]
        xs=[r['income']['percentage']for r in eligible]
        subjects[subject]=dict(applicable=sum(r['subjects']['ES'][subject]['applicable']for r in rows),usable=len(eligible),
            prospective_nces_school_ids=[r['school_id']for r in eligible],prospective_native_school_ids=[r['native_school_id']for r in eligible],
            exclusions=dict(sorted(Counter(r['subjects']['ES'][subject]['exclusion']for r in rows if r['subjects']['ES'][subject]['exclusion']).items())),
            income=dict(minimum=min(xs)if xs else None,maximum=max(xs)if xs else None,distinct=len(set(xs)),sample_standard_deviation=statistics.stdev(xs)if len(xs)>1 else None),
            charter=dict(sorted(Counter(r['ccd_flags']['charter']for r in eligible).items())),
            school_type=dict(sorted(Counter(r['ccd_flags']['school_type']for r in eligible).items())),
            prospective_30_school_floor_met=len(eligible)>=30,approved_for_modeling=False)
    return dict(profiles=len(rows),nces_school_ids=[r['school_id']for r in rows],native_school_ids=[r['native_school_id']for r in rows],subjects=subjects,approved_for_source=False,approved_for_modeling=False)

def original_identity_fingerprint(inputs):
    result={}
    for table in ORIGINAL_INVENTORIES:
        result[table]=[dict(source_row=e['source_row'],identity=[e['raw'].get(k)for k in (['SCHOOL_YEAR','LEAID','ST_LEAID','NCESSCH','ST_SCHID','GRADE','TOTAL_INDICATOR']if table in{'directory','membership','lea_directory','lea_membership'}else['Academic Year','County Code','District Code','School Code','Type ID','Test Year','Test Type','Test ID','Student Group ID','Grade'])])for e in inputs[table]['records']]
    return fingerprint(result)

def validate_definition_references(definitions):
    if set(definitions)!={'frpm','research_layout','ccd_state_notes','research_page'}:raise ValueError('Complete original primary definitions required')
    for key in ['frpm','research_layout','ccd_state_notes']:
        block=definitions[key]
        if block['source_id']!=key or not block['worksheets']:raise ValueError('Original primary workbook evidence required')
        for sheet in block['worksheets']:
            if type(sheet['worksheet'])is not str or not sheet['rows']:raise ValueError('Original worksheet namespace required')
            previous=0
            for e in sheet['rows']:
                if type(e['source_row'])is not int or e['source_row']<=previous or type(e['raw'])is not list:raise ValueError('Strict original primary worksheet row references required')
                previous=e['source_row']
    page=definitions['research_page']
    if page['source_id']!='research_page'or page['encoding']!='utf-8'or not page['lines']:raise ValueError('Original primary research HTML required')
    if any(type(e['source_line'])is not int or e['source_line']!=i or type(e['html'])is not str for i,e in enumerate(page['lines'],1)):raise ValueError('Exact integer primary HTML line references required')
    text=' '.join(e['html']for e in page['lines'])
    if 'fewer than 11 students'not in text or '2024&ndash;25 Research Files'not in text:raise ValueError('Original2025 suppression/population definitions changed')

def read_ccd_rows(key):
    path = ROOT / SOURCES[key]['path']
    member = path.stem + '.csv'
    process = subprocess.Popen(['unzip', '-p', str(path), member], stdout=subprocess.PIPE,
                               text=True, encoding='utf-8-sig')
    records, omitted, total, cells, populations = [], [], 0, 0, Counter()
    digest = hashlib.sha256(b'[')
    try:
        parser = csv.DictReader(process.stdout)
        for ordinal, row in enumerate(parser, 2):
            if row['LEAID'] != LEA and row['ST_LEAID'] != NATIVE_LEA:
                continue
            total += 1
            if 'directory' in key or row['TOTAL_INDICATOR'] in {'Education Unit Total', 'Subtotal 4 - By Grade'}:
                records.append(dict(source_row=ordinal, raw=row))
            else:
                if omitted:
                    digest.update(b',')
                digest.update(json.dumps(dict(source_row=ordinal, raw=row), sort_keys=True, separators=(',', ':'), ensure_ascii=False).encode())
                identity = ['SCHOOL_YEAR', 'ST', 'LEAID', 'ST_LEAID', 'ST_SCHID', 'NCESSCH', 'GRADE', 'RACE_ETHNICITY', 'SEX', 'TOTAL_INDICATOR', 'DMS_FLAG']
                omitted.append(dict(source_row=ordinal, identity_population={k: row[k] for k in identity if k in row}))
                cells += len(row)
                populations[row.get('TOTAL_INDICATOR', 'not_applicable')] += 1
        if process.wait():
            raise ValueError('Original CCD stream failed')
        digest.update(b']')
        return dict(source_id=key, member=member, header_source_row=1, header=parser.fieldnames, records=records,
                    complete_district_inventory=dict(original_district_rows=total, retained_rows=len(records),
                        omitted_rows=len(omitted), omitted_source_ordinals=[r['source_row'] for r in omitted],
                        omitted_identity_population_inventory=omitted, omitted_typed_cells=cells,
                        omitted_raw_ordinal_sha256=digest.hexdigest(), omitted_populations=dict(sorted(populations.items()))))
    finally:
        process.stdout.close()
        if process.poll() is None:
            process.terminate()
            process.wait()

def table_inventory(source_id, header, records, omitted, **extra):
    identity = ['County Code', 'District Code', 'School Code', 'Type ID', 'Test Year', 'Test Type', 'Test ID', 'Student Group ID', 'Grade']
    inventory = [dict(source_row=e['source_row'], identity_population={k: e['raw'][k] for k in identity if k in e['raw']}) for e in omitted]
    return dict(source_id=source_id, header=header, records=records,
                complete_district_inventory=dict(original_district_rows=len(records)+len(omitted), retained_rows=len(records),
                    omitted_rows=len(omitted), omitted_source_ordinals=[e['source_row'] for e in omitted],
                    omitted_identity_population_inventory=inventory, omitted_typed_cells=sum(len(e['raw']) for e in omitted),
                    omitted_raw_ordinal_sha256=fingerprint(omitted),
                    omitted_populations=dict(sorted(Counter(e['raw'].get('Type ID', 'not_applicable') for e in omitted).items()))), **extra)

def workbook_definition(source_id, selected=None):
    workbook = openpyxl.load_workbook(ROOT / SOURCES[source_id]['path'], read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook:
            if selected is not None and sheet.title not in selected:
                continue
            rows = [dict(source_row=i, raw=[native_cell(v) for v in values]) for i, values in enumerate(sheet.values, 1)]
            sheets.append(dict(worksheet=sheet.title, rows=rows))
    finally:
        workbook.close()
    return dict(source_id=source_id, worksheets=sheets)

def extract_raw_inputs():
    require_current_originals()
    for key, pin in SOURCES.items():
        path = ROOT / pin['path']
        assert path.stat().st_size == pin['bytes'] and hashlib.sha256(path.read_bytes()).hexdigest() == pin['sha256'], key
    result = {key: read_ccd_rows(key) for key in ['directory', 'membership', 'lea_directory', 'lea_membership']}
    workbook = openpyxl.load_workbook(ROOT / SOURCES['frpm']['path'], read_only=True, data_only=True)
    sheet = workbook['FRPM School-Level Data']
    it = iter(sheet.values)
    title, original_header = list(next(it)), list(next(it))
    header = [' '.join(str(v).split()) for v in original_header]
    rows = []
    for i, values in enumerate(it, 3):
        raw = dict(zip(header, [native_cell(v) for v in values]))
        if raw['County Code'] + raw['District Code'] == CDS_DISTRICT:
            rows.append(dict(source_row=i, raw=raw))
    workbook.close()
    result['frpm'] = table_inventory('frpm', header, rows, [], worksheet='FRPM School-Level Data',
                                   title_source_row=1, title=title, header_source_row=2, original_header=original_header)
    with zipfile.ZipFile(ROOT / SOURCES['assessment']['path']) as archive:
        for key, member in [('assessment', 'sb_ca2025_1_csv_v1.txt'), ('entities', 'sb_ca2025entities_csv.txt')]:
            with archive.open(member) as stream:
                reader = csv.DictReader(io.TextIOWrapper(stream, encoding='cp1252', newline=''), delimiter='^')
                rows, omitted = [], []
                for i, raw in enumerate(reader, 2):
                    if raw['County Code'] + raw['District Code'] != CDS_DISTRICT:
                        continue
                    if set(raw) != set(reader.fieldnames):
                        raise ValueError('Malformed complete original CA row')
                    entry = dict(source_row=i, raw=raw)
                    if raw['Type ID'] in {'7', '9', '10'}:
                        rows.append(entry)
                    else:
                        omitted.append(entry)
                result[key] = table_inventory('assessment', reader.fieldnames, rows, omitted, member=member,
                                             header_source_row=1, original_ordinal_kind='nonempty_CSV_record_ordinal_including_header')
    definitions = dict(
        frpm=workbook_definition('frpm', {'Title Page', 'Data Field Descriptions'}),
        research_layout=workbook_definition('research_layout'))
    notes = openpyxl.load_workbook(ROOT / SOURCES['ccd_state_notes']['path'], read_only=True, data_only=True)
    sheets = []
    for sheet in notes:
        rows = [dict(source_row=i, raw=[native_cell(v) for v in values]) for i, values in enumerate(sheet.values, 1)
                if i <= 3 or any(isinstance(v, str) and ('California' in v or 'San Diego' in v or v == 'CALIFORNIA') for v in values)]
        if any(row['source_row'] > 3 for row in rows):
            sheets.append(dict(worksheet=sheet.title, rows=rows))
    notes.close()
    definitions['ccd_state_notes'] = dict(source_id='ccd_state_notes', worksheets=sheets)
    definitions['research_page'] = dict(source_id='research_page', encoding='utf-8',
        lines=[dict(source_line=i, html=line) for i, line in enumerate((ROOT / SOURCES['research_page']['path']).read_text().splitlines(), 1)])
    result['definitions'] = definitions
    planning = json.loads((ROOT / 'data/source/district-comparison-candidates.json').read_text())
    agency = next(a for a in planning['agencies'] if a['nces_lea_id'] == LEA)
    fields = ['nces_lea_id', 'state', 'native_lea_id', 'enrollment', 'enrollment_flag', 'school_directory_operational_count',
              'potential_cohorts', 'potential_school_ids', 'directory_raw', 'directory_source_row', 'membership_raw', 'membership_source_row']
    result['planning'] = dict(source='data/source/district-comparison-candidates.json', sources=planning['sources'],
                              record={k: agency[k] for k in fields})
    return result

def derive(document):
    require_current_originals()
    base=dict(schema_version=1,state='CA',year=TEST_YEAR,academic_year=YEAR,nces_lea_id=LEA,native_lea_id=NATIVE_LEA,native_district_id=CDS_DISTRICT,status='audit_pending',scope='source_cohort_audit_only',approved_for_source=False,approved_for_modeling=False,sources=SOURCES,policy=POLICY)
    if any(fingerprint(document.get(k))!=fingerprint(v)for k,v in base.items()):raise ValueError('San Diego source-only immutable metadata/flags changed')
    inputs=document['raw_inputs']
    if fingerprint(inputs)!=RAW_INPUT_SHA:raise ValueError('Complete original retained/omitted San Diego inputs changed')
    metadata={k:{field:value for field,value in inputs[k].items()if field!='records'}for k in ORIGINAL_INVENTORIES}
    if fingerprint(metadata)!=ORIGINAL_SCHEMA_SHA or fingerprint(inputs['definitions'])!=DEFINITION_INPUT_SHA or original_identity_fingerprint(inputs)!=ORIGINAL_IDENTITY_SHA:raise ValueError('Original schema/definitions/strict identity projection changed')
    validate_definition_references(inputs['definitions'])
    for table,size in ORIGINAL_INVENTORIES.items():
        block=inputs[table]
        if len(block['records'])!=size or fingerprint(metadata[table])!=ORIGINAL_METADATA[table]or len(set(block['header']))!=len(block['header']):raise ValueError('Complete original needed/omitted schema inventory changed')
        previous=0
        for e in block['records']:
            if type(e['source_row'])is not int or e['source_row']<=previous or type(e['raw'])is not dict or set(e['raw'])!=set(block['header']):raise ValueError('Complete strict typed original row/header required')
            if table!='frpm'and any(type(v)is not str for v in e['raw'].values()):raise ValueError('CSV raw native fields must preserve strings')
            previous=e['source_row']
        if type(block['header_source_row'])is not int or block['header_source_row']!=(2 if table=='frpm'else 1):raise ValueError('Strict original header source-row identity required')
    directory={};native_to_nces={};totals={};grades=defaultdict(dict)
    for e in inputs['directory']['records']:
        n,key=ccd_identity(e['raw'])
        if n in directory or key in native_to_nces:raise ValueError('Duplicate original exact CCD identity')
        directory[n]=e;native_to_nces[key]=n
    for e in inputs['membership']['records']:
        n,key=ccd_identity(e['raw']);raw=e['raw']
        if n not in directory or native_to_nces[key]!=n:raise ValueError('CCD membership without exact same-year roster identity')
        if raw['TOTAL_INDICATOR']=='Education Unit Total':
            if n in totals:raise ValueError('Duplicate original CCD membership total')
            totals[n]=e
        elif raw['TOTAL_INDICATOR']=='Subtotal 4 - By Grade':
            if raw['GRADE']in grades[n]:raise ValueError('Duplicate original CCD membership grade')
            grades[n][raw['GRADE']]=e
        else:raise ValueError('Unexpected retained CCD demographic population')
    native={};assessments=defaultdict(list);entities={}
    for e in inputs['frpm']['records']:
        key=cds(e['raw'])
        if key in native:raise ValueError('Duplicate original native FRPM CDS')
        income_evidence(e);native[key]=e
    seen=set()
    for e in inputs['assessment']['records']:
        raw=e['raw'];key=cds(raw);identity=(key,raw['Test ID'],raw['Grade'])
        if (raw['Type ID']not in SCHOOL_TYPES or raw['Test Year']!=str(TEST_YEAR)or raw['Test Type']!='B'or raw['Student Group ID']!='1'or raw['Test ID']not in SUBJECTS.values()or raw['Grade']not in{'3','4','5','6','7','8','11','13'}or identity in seen):raise ValueError('Original school All Students SB subject/grade/year population changed')
        seen.add(identity);assessments[(key,raw['Test ID'])].append(e)
    for e in inputs['entities']['records']:
        raw=e['raw'];key=cds(raw)
        if raw['Type ID']not in SCHOOL_TYPES or raw['Test Year']!=str(TEST_YEAR)or key in entities:raise ValueError('Duplicate/mis-scoped original entity record')
        entities[key]=e
    lea_directory=inputs['lea_directory']['records']
    if len(lea_directory)!=1:raise ValueError('Exact original agency directory required')
    ccd_identity(lea_directory[0]['raw'],False)
    for e in inputs['lea_membership']['records']:ccd_identity(e['raw'],False)
    lea_total=[e for e in inputs['lea_membership']['records']if e['raw']['TOTAL_INDICATOR']=='Education Unit Total']
    if len(lea_total)!=1:raise ValueError('Exact original agency total required')
    planning=inputs['planning']['record']
    if fingerprint([planning['nces_lea_id'],planning['native_lea_id'],planning['state']])!=fingerprint([LEA,NATIVE_LEA,'CA']):raise ValueError('Planning exact agency identity changed')
    for field in ['directory_source_row','membership_source_row','enrollment','school_directory_operational_count']:
        if type(planning[field])is not int or planning[field]<=0:raise ValueError('Planning reference/count must be exact positive integer')
    if planning['enrollment']!=ccd_count(lea_total[0]['raw'])or planning['membership_raw']!=lea_total[0]['raw']or planning['directory_source_row']!=lea_directory[0]['source_row']or planning['membership_source_row']!=lea_total[0]['source_row']:raise ValueError('Planning primary agency membership replay changed')
    links=[];outside=[];differences=[];bounds=[]
    for n,e in directory.items():
        raw=e['raw'];key=ccd_identity(raw)[1];profile=native.get(key);operational=raw['UPDATED_STATUS']in OPERATIONAL
        member=enrollment_evidence(totals.get(n),grades[n],raw)
        row=dict(school_id=n,native_school_id=key,name=raw['SCH_NAME'],ccd_source_row=e['source_row'],operational=operational,
            ccd_flags=dict(status=raw['UPDATED_STATUS'],status_text=raw['UPDATED_STATUS_TEXT'],charter=raw['CHARTER_TEXT'],school_type=raw['SCH_TYPE_TEXT'],level=raw['LEVEL']),
            ccd_enrollment=member,native_profile=None if profile is None else dict(source_row=profile['source_row'],raw=profile['raw']),matched=profile is not None)
        if not profile:
            row.update(exclusion='No exact same-year native FRPM profile',prospective_lower_profile=False);links.append(row);continue
        p=profile['raw'];offered=offered_evidence(raw,p);income=income_evidence(profile)
        charter=p['Charter School (Y/N)'].strip()
        if charter not in{'Yes','No'}or charter!=raw['CHARTER_TEXT']:raise ValueError('Original CCD/native charter labels disagree')
        typed={v['raw']['Type ID']for tid in SUBJECTS.values()for v in assessments[(key,tid)]}
        if charter=='No'and typed-{'7'}:raise ValueError('Noncharter native school outcome type changed')
        prospect=operational and offered['pure_lower_candidate']and not member['high_positive']and not member['high_unknown']
        if not offered['literal_bounds_agree']:bounds.append(dict(school_id=n,native_school_id=key,evidence=offered))
        if member['total']is not None and income['enrollment']is not None and member['total']!=income['enrollment']:differences.append(dict(school_id=n,native_school_id=key,ccd_enrollment=member['total'],native_K12_enrollment=income['enrollment'],native_minus_ccd=income['enrollment']-member['total']))
        row.update(income=income,grade_evidence=offered,prospective_lower_profile=prospect,
            original_entity=None if key not in entities else entities[key],
            native_administrative_labels={k:p[k]for k in ['School Type','Educational Option Type','Charter School (Y/N)','Charter Funding Type','Independently Reporting Charter (IRC)']},
            programs_admissions_virtual_provider='Unknown; no attribution from school name',subjects={})
        for level,expected in [('ES',offered['expected_tested_grades']),('HS',[11]if grade(p['Low Grade'])is not None and grade(p['High Grade'])is not None and grade(p['Low Grade'])<=11<=grade(p['High Grade'])else[])]:
            values={s:subject_evidence(assessments[(key,tid)],expected,income)for s,tid in SUBJECTS.items()}
            m,r=values['math'],values['reading'];usable=m['source_usable']and r['source_usable']
            values['combined']=dict(applicable=m['applicable']and r['applicable'],source_usable=usable,
                proficiency=(m['proficiency']+r['proficiency'])/2 if usable else None,valid_scores=None,
                valid_score_components={'math':m['valid_scores'],'reading':r['valid_scores']},
                exclusion=None if usable else'Requires usable original Math and ELA: '+'; '.join(s+': '+v['exclusion']for s,v in [('math',m),('ELA',r)]if not v['source_usable']))
            row['subjects'][level]=values
        row['exclusion']=None if prospect else offered['exclusion']or'Outside pure lower operational population'
        links.append(row)
    for key,e in native.items():
        if key not in native_to_nces:
            outside.append(dict(native_school_id=key,profile=e,entity=entities.get(key),assessment_source_rows=[v['source_row']for tid in SUBJECTS.values()for v in assessments[(key,tid)]],exclusion='Native reporting prefix without exact CCD LEA roster membership; never a district model member'))
    selected=[r for r in links if r.get('prospective_lower_profile')]
    known_totals=[r['ccd_enrollment']['total']for r in links if r['operational']and r['ccd_enrollment']['total']is not None]
    native_totals=[r['income']['enrollment']for r in links if r.get('income')and r['income']['enrollment']is not None]
    expected_planning={level:set(planning['potential_school_ids'][level])for level in ['ES','HS']}
    if {level:len(ids)for level,ids in expected_planning.items()}!=planning['potential_cohorts']:raise ValueError('Planning complete cohort-ID counts changed')
    return dict(roster_links=links,outside_roster_native_profiles=outside,
        grade_bound_differences=bounds,enrollment_reconciliation=dict(ccd_LEA_enrollment=ccd_count(lea_total[0]['raw']),
            known_operational_school_enrollment=sum(known_totals),unknown_operational_totals=[r['school_id']for r in links if r['operational']and r['ccd_enrollment']['total']is None],
            ccd_LEA_minus_known_school_sum=ccd_count(lea_total[0]['raw'])-sum(known_totals),native_exact_roster_K12_enrollment=sum(native_totals),
            ccd_LEA_minus_native_K12_sum=ccd_count(lea_total[0]['raw'])-sum(native_totals),individual_known_differences=differences,
            note='Different snapshots/populations; no unmatched difference is assigned to a school or treated as zero.'),
        prospective_cohorts={'pure_grade_schools':summarize(links,lambda r:r.get('prospective_lower_profile'))},
        planning_reconciliation=dict(discovery_ES_ids=sorted(expected_planning['ES']),source_ES_ids=[r['school_id']for r in selected],
            planning_only=sorted(expected_planning['ES']-{r['school_id']for r in selected}),source_only=sorted({r['school_id']for r in selected}-expected_planning['ES']),planning_never_approves_model=True),
        coverage=dict(directory_records=len(links),operational_schools=sum(r['operational']for r in links),exact_native_matches=sum(r['matched']for r in links),
            outside_reporting_prefix_profiles=len(outside),pure_lower_profiles=len(selected),
            source_usable_math=sum(r['subjects']['ES']['math']['source_usable']for r in selected),source_usable_reading=sum(r['subjects']['ES']['reading']['source_usable']for r in selected),
            source_usable_combined=sum(r['subjects']['ES']['combined']['source_usable']for r in selected),
            substantive_grade_scope_holds=[r['school_id']for r in links if r.get('grade_evidence')and not r['grade_evidence']['configuration_agreement']],
            roster_charter=dict(sorted(Counter(r['ccd_flags']['charter']for r in links).items())),
            lower_school_types=dict(sorted(Counter(r['ccd_flags']['school_type']for r in selected).items())),
            full_needed_original_records=sum(ORIGINAL_INVENTORIES.values()),full_needed_original_typed_cells=sum(len(e['raw'])for table in ORIGINAL_INVENTORIES for e in inputs[table]['records']),
            historical_approvals_false=True))

def validate(document):
    expected=derive(document)
    base={'schema_version','state','year','academic_year','nces_lea_id','native_lea_id','native_district_id','status','scope','approved_for_source','approved_for_modeling','sources','policy','raw_inputs'}
    if set(document)!=base|set(expected)or any(fingerprint(document.get(k))!=fingerprint(v)for k,v in expected.items()):raise ValueError('San Diego exact derived original evidence/cohort/ratios/holds changed')
    return expected['coverage']

def build(inputs,output=OUTPUT):
    document=dict(schema_version=1,state='CA',year=TEST_YEAR,academic_year=YEAR,nces_lea_id=LEA,native_lea_id=NATIVE_LEA,native_district_id=CDS_DISTRICT,status='audit_pending',scope='source_cohort_audit_only',approved_for_source=False,approved_for_modeling=False,sources=SOURCES,policy=POLICY,raw_inputs=inputs)
    document.update(derive(document));validate(document)
    Path(output).write_text(json.dumps(document,sort_keys=True,separators=(',',':'),ensure_ascii=False,allow_nan=False)+'\n')
    return document['coverage']

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--extract',action='store_true',help='Reopen all pinned whole original CCD/CDE tables and primary definitions')
    parser.add_argument('--output',type=Path,default=OUTPUT)
    args=parser.parse_args()
    result=build(extract_raw_inputs(),args.output)if args.extract else validate(json.loads(args.output.read_text()))
    print(json.dumps(result,sort_keys=True))
